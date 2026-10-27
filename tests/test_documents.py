"""Reading any file type: documents.extract and the read_file / view_image tools."""

from __future__ import annotations

import json
import shutil
import sqlite3
import zipfile

import pytest

from kestrel.documents import extract, parse_pages, to_viewable
from kestrel.tools import Workspace, builtin_tools

DOCX_XML = """<?xml version="1.0" encoding="UTF-8"?>
<w:document xmlns:w="http://schemas.openxmlformats.org/wordprocessingml/2006/main"><w:body>
<w:p><w:pPr><w:pStyle w:val="Heading1"/></w:pPr><w:r><w:t>Quarterly report</w:t></w:r></w:p>
<w:p><w:r><w:t>Revenue grew </w:t></w:r><w:r><w:t>12%.</w:t></w:r></w:p>
<w:p><w:pPr><w:numPr/></w:pPr><w:r><w:t>Hire two engineers</w:t></w:r></w:p>
<w:tbl><w:tr><w:tc><w:p><w:r><w:t>Region</w:t></w:r></w:p></w:tc><w:tc><w:p><w:r><w:t>Sales</w:t></w:r></w:p></w:tc></w:tr>
<w:tr><w:tc><w:p><w:r><w:t>EU</w:t></w:r></w:p></w:tc><w:tc><w:p><w:r><w:t>40</w:t></w:r></w:p></w:tc></w:tr></w:tbl>
</w:body></w:document>"""

SLIDE_XML = """<?xml version="1.0"?><p:sld xmlns:p="http://schemas.openxmlformats.org/presentationml/2006/main"
xmlns:a="http://schemas.openxmlformats.org/drawingml/2006/main"><p:cSld><p:spTree><p:sp><p:txBody>
<a:p><a:r><a:t>{title}</a:t></a:r></a:p><a:p><a:r><a:t>{body}</a:t></a:r></a:p></p:txBody></p:sp></p:spTree></p:cSld></p:sld>"""


def minimal_pdf(lines: list[str]) -> bytes:
    """A real one-page PDF with a text layer (Helvetica), built by hand so the test needs no PDF writer."""
    content = "BT /F1 18 Tf 72 720 Td " + " ".join(f"({t}) Tj 0 -24 Td" for t in lines) + " ET"
    objs = ["<< /Type /Catalog /Pages 2 0 R >>", "<< /Type /Pages /Kids [3 0 R] /Count 1 >>",
            "<< /Type /Page /Parent 2 0 R /MediaBox [0 0 612 792] /Contents 4 0 R /Resources << /Font << /F1 5 0 R >> >> >>",
            f"<< /Length {len(content)} >>\nstream\n{content}\nendstream",
            "<< /Type /Font /Subtype /Type1 /BaseFont /Helvetica >>"]
    out, offsets = b"%PDF-1.4\n", []
    for i, o in enumerate(objs, 1):
        offsets.append(len(out))
        out += f"{i} 0 obj\n{o}\nendobj\n".encode()
    xref = len(out)
    out += f"xref\n0 {len(objs) + 1}\n0000000000 65535 f \n".encode()
    out += b"".join(f"{o:010d} 00000 n \n".encode() for o in offsets)
    out += f"trailer << /Size {len(objs) + 1} /Root 1 0 R >>\nstartxref\n{xref}\n%%EOF".encode()
    return out


@pytest.fixture
def files(tmp_path):
    d = tmp_path
    with zipfile.ZipFile(d / "report.docx", "w") as z:
        z.writestr("word/document.xml", DOCX_XML)
    with zipfile.ZipFile(d / "deck.pptx", "w") as z:
        z.writestr("ppt/slides/slide1.xml", SLIDE_XML.format(title="Roadmap", body="Ship v1 in May"))
        z.writestr("ppt/slides/slide2.xml", SLIDE_XML.format(title="Risks", body="Hiring"))
    import openpyxl
    wb = openpyxl.Workbook()
    ws = wb.active
    ws.title = "Sales"
    ws.append(["region", "units", "price"])
    for r in [("EU", 40, 9.5), ("US", 55, 10.0), ("APAC", 12, 8.0)]:
        ws.append(r)
    wb.create_sheet("Notes").append(["checked by", "Ada"])
    wb.save(d / "sales.xlsx")
    (d / "data.csv").write_text("name,age,city\nAda,36,London\nLinus,28,Helsinki\nGrace,45,Arlington\n")
    (d / "conf.json").write_text(json.dumps({"service": {"port": 8080, "hosts": ["a", "b"]}, "debug": False}))
    (d / "people.json").write_text(json.dumps([{"name": "Ada", "lang": "en"}, {"name": "Yuki", "lang": "ja"}]))
    (d / "schema.jsonld").write_text(json.dumps({"@context": "https://schema.org", "@type": "Person", "name": "Ada"}))
    (d / "ci.yml").write_text("jobs:\n  test:\n    runs-on: ubuntu-latest\n    steps: [checkout, test]\n")
    (d / "app.tsx").write_text("export const App = () => <div>hi</div>;\n")
    (d / "page.html").write_text("<html><body><h1>Title</h1><script>evil()</script><p>Hello &amp; welcome</p></body></html>")
    (d / "notes.rtf").write_text(r"{\rtf1\ansi{\fonttbl\f0 Arial;}\f0 Hello \b bold\b0  world\par Second line}")
    (d / "nb.ipynb").write_text(json.dumps({"cells": [
        {"cell_type": "markdown", "source": ["# Analysis"]},
        {"cell_type": "code", "source": ["print(1+1)"], "outputs": [{"text": ["2\n"]}]}]}))
    (d / "scan.pdf").write_bytes(minimal_pdf(["Invoice 1042", "Total due 99 EUR"]))
    con = sqlite3.connect(d / "app.db")
    con.execute("create table users (id integer primary key, email text)")
    con.executemany("insert into users (email) values (?)", [("a@x.io",), ("b@x.io",)])
    con.commit()
    con.close()
    with zipfile.ZipFile(d / "bundle.zip", "w") as z:
        z.writestr("src/main.py", "print('hi')")
        z.writestr("README.md", "# bundle")
    from PIL import Image
    Image.new("RGB", (40, 20), "red").save(d / "logo.bmp")
    (d / "blob.bin").write_bytes(b"\x00\x01KESTRELMAGIC\x02" + bytes(range(256)))
    return d


async def test_documents_become_text_and_tables(files):
    docx = await extract(files / "report.docx")
    assert "# Quarterly report" in docx.text and "Revenue grew 12%." in docx.text and "- Hire two engineers" in docx.text
    assert docx.tables[0].columns == ["Region", "Sales"] and "| EU | 40 |" in docx.text
    pptx = await extract(files / "deck.pptx")
    assert pptx.pages == 2 and "## Slide 2\nRisks\nHiring" in pptx.text
    pdf = await extract(files / "scan.pdf")
    assert pdf.kind == "pdf" and pdf.pages == 1 and "Invoice 1042" in pdf.text and "Total due 99 EUR" in pdf.text
    rtf = await extract(files / "notes.rtf")
    assert "Hello bold world" in rtf.text and "Second line" in rtf.text
    nb = await extract(files / "nb.ipynb")
    assert "# Analysis" in nb.text and "print(1+1)" in nb.text and "2" in nb.text


async def test_structured_data_is_outlined_and_summarised(files):
    xlsx = await extract(files / "sales.xlsx")
    assert xlsx.meta["sheets"] == ["Sales", "Notes"]
    sales = xlsx.tables[0]
    assert sales.columns == ["region", "units", "price"] and sales.total_rows == 3
    assert "units: numeric, min 12, max 55" in xlsx.text
    only = await extract(files / "sales.xlsx", sheet="Notes")
    assert [t.name for t in only.tables] == ["Notes"]
    csv_ = await extract(files / "data.csv")
    assert csv_.tables[0].columns == ["name", "age", "city"] and csv_.tables[0].total_rows == 3
    assert "age: numeric, min 28, max 45" in csv_.text
    conf = await extract(files / "conf.json")
    assert "service: object (2 keys)" in conf.text and "hosts: array (2)" in conf.text
    people = await extract(files / "people.json")
    assert people.kind == "table" and people.tables[0].columns == ["name", "lang"]
    ld = await extract(files / "schema.jsonld")
    assert "@type" in ld.text
    yml = await extract(files / "ci.yml")
    assert "jobs: object (1 keys)" in yml.text and "runs-on: ubuntu-latest" in yml.text


async def test_containers_images_and_binaries(files):
    db = await extract(files / "app.db")
    assert db.kind == "database" and "users (2 rows)" in db.text and "a@x.io" in db.text
    z = await extract(files / "bundle.zip")
    assert "src/main.py" in z.text and z.tables[0].total_rows == 2
    img = await extract(files / "logo.bmp")
    assert img.kind == "image" and img.meta["width"] == 40
    blob = await extract(files / "blob.bin")
    assert blob.kind == "binary" and "KESTRELMAGIC" in blob.text
    page = await extract(files / "page.html")
    assert "evil()" not in page.meta["readable"] and "Hello & welcome" in page.meta["readable"]


async def test_read_file_tool_reads_anything_and_view_image_converts(files):
    seen = []

    class FakeVision:
        async def complete(self, messages, **kw):
            from kestrel.llm import Completion
            seen.append(messages[-1]["content"][1]["image_url"]["url"][:22])
            return Completion(content="a red rectangle")

    tools = {t.name: t for t in builtin_tools(Workspace(files), allow_shell=False, output_limit=20000, vision=FakeVision())}
    read = tools["read_file"].handler
    assert "[document: docx]" in await read({"path": "report.docx"})
    assert "Invoice 1042" in await read({"path": "scan.pdf", "pages": "1"})
    plain = await read({"path": "scan.pdf"})  # this PDF is pure ASCII: it must still be extracted, not shown as source
    assert "[pdf: pdf, 1 pages]" in plain and "%PDF" not in plain
    assert "    1\texport const App" in await read({"path": "app.tsx"})          # source stays line-numbered
    assert "1\tname,age,city" in await read({"path": "data.csv"})                # raw by default, for editing
    assert "age: numeric" in await read({"path": "data.csv", "structured": True})  # or summarised
    assert "region" in await read({"path": "sales.xlsx", "sheet": "Sales"})
    assert await tools["view_image"].handler({"path": "logo.bmp"}) == "a red rectangle"
    assert seen == ["data:image/png;base64,"]  # BMP converted to PNG for the vision model


@pytest.mark.skipif(not shutil.which("pdftoppm"), reason="needs poppler-utils")
async def test_pdf_pages_render_for_the_vision_model(files):
    png = await to_viewable(files / "scan.pdf", 1)
    assert png.suffix == ".png" and png.stat().st_size > 1000


def test_page_ranges():
    assert parse_pages("2-4,7", 8) == [2, 3, 4, 7]
    assert parse_pages(None, 3) == [1, 2, 3]
    assert parse_pages("9", 3) == []
